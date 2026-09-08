"""Import and export functionality for Google Sheets / Excel workbook (Invest.xlsx).

Implements:
- Transactional, idempotent-safe import from XLSX (§1.1, §1.2)
- Handling of multi-buy aggregation (§2.2)
- Preservation and reporting of manual/special entries (per user instruction)
- Round-trip export back to Excel format (§1.3)
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .db import get_connection
from .db.backup import create_backup
from .db.repositories import dividends as div_repo
from .db.repositories import holdings as hold_repo
from .db.repositories import summary as sum_repo
from .portfolio_tracker import (
    compute_buy_charge,
    compute_sell_charge,
    compute_summary_panel,
    compute_totals,
    enrich_holding_row,
    enrich_sold_row,
)
from .sheet_format import (
    DIVIDEND_HEADER_LABEL,
    DIVIDEND_SECTION_LABEL,
    FALLBACK_LAYOUT,
    HOLDING_COLUMNS,
    HOLDINGS_HEADER_LABEL,
    SHEET_NAMES,
    SOLD_SECTION_LABELS,
    SOLD_START_COLUMN,
    SUMMARY_LABEL_COLUMN,
    SUMMARY_PANEL_CONFIG,
    SUMMARY_VALUE_COLUMN,
    TOTAL_ROW_LABEL,
    format_sheet_date,
    map_app_to_person,
    normalize_label,
    parse_sheet_date,
)

logger = logging.getLogger(__name__)


def _safe_number(val: Any, default: float = 0.0) -> tuple[float, bool]:
    """Safely convert cell value to float.

    Returns (value, is_unusual_flag).
    """
    if val is None:
        return default, False
    if isinstance(val, (int, float)):
        return float(val), False
    s = str(val).strip()
    if not s or s == "-":
        return default, False
    try:
        clean = s.replace(",", "").replace("₹", "").strip()
        return float(clean), False
    except ValueError:
        return default, True


def _is_manual_charge(cell_value: Any, is_formula: bool, expected: float) -> bool:
    """Decide whether a charge cell holds a user override rather than a derived value.

    Excel stores computed charges as formulas, but §1.3 requires the exporter to
    write them out as plain values -- so a literal that matches what the app would
    compute is treated as derived, keeping the round-trip idempotent.
    """
    if is_formula or cell_value is None or cell_value < 0:
        return False
    return abs(cell_value - expected) > 0.01


# ---------------------------------------------------------------------------
# Layout discovery (§1.2)
# ---------------------------------------------------------------------------
# Section positions are located by their marker labels rather than fixed row
# numbers. An exported workbook shifts every section as the data grows, so
# hard-coded rows would break the §1.3 round-trip guarantee.

_MARKER_SEARCH_LIMIT = 40


def _label_at(ws: Any, row: int, col: int) -> str:
    return normalize_label(ws.cell(row, col).value)


def _find_label_row(
    ws: Any,
    col: int,
    labels: tuple[str, ...],
    start: int = 1,
    end: int | None = None,
) -> int | None:
    """Return the first row in ``col`` whose text matches one of ``labels``."""
    last = min(end or ws.max_row, ws.max_row)
    for r in range(start, last + 1):
        if _label_at(ws, r, col) in labels:
            return r
    return None


def _resolve_layout(ws: Any, portfolio: str) -> dict[str, int]:
    """Locate the Holdings, Sold and Dividend blocks on a sheet.

    Falls back to the original ``Invest.xlsx`` row numbers when a marker is absent.
    """
    fallback = FALLBACK_LAYOUT[portfolio]
    sold_col = SOLD_START_COLUMN[portfolio]

    header_row = _find_label_row(ws, 1, (HOLDINGS_HEADER_LABEL,), 1, _MARKER_SEARCH_LIMIT)
    holdings_start = header_row + 1 if header_row else fallback["holdings_start"]

    if sold_col == 1:
        # Stacked layout: the Sold block follows a standalone "Sold" marker.
        marker = _find_label_row(ws, 1, SOLD_SECTION_LABELS, holdings_start)
        if marker:
            sold_start = marker + 1
            if _label_at(ws, sold_start, 1) == HOLDINGS_HEADER_LABEL:
                sold_start += 1
        else:
            sold_start = fallback["sold_start"]
    else:
        sold_header = _find_label_row(ws, sold_col, (HOLDINGS_HEADER_LABEL,), 1, _MARKER_SEARCH_LIMIT)
        sold_start = sold_header + 1 if sold_header else fallback["sold_start"]

    div_marker = _find_label_row(ws, 1, (DIVIDEND_SECTION_LABEL,), holdings_start)
    if div_marker:
        dividend_start = div_marker + 1
        if _label_at(ws, dividend_start, 1) == DIVIDEND_HEADER_LABEL:
            dividend_start += 1
    else:
        dividend_start = fallback["dividend_start"]

    return {
        "holdings_start": holdings_start,
        "sold_col": sold_col,
        "sold_start": sold_start,
        "dividend_start": dividend_start,
    }


# ---------------------------------------------------------------------------
# Importer (§1.1, §1.2)
# ---------------------------------------------------------------------------

def check_existing_data() -> dict[str, int]:
    """Check counts of existing portfolio tracker data."""
    conn = get_connection()
    h_count = conn.execute("SELECT COUNT(*) AS c FROM holding").fetchone()["c"]
    s_count = conn.execute("SELECT COUNT(*) AS c FROM sale").fetchone()["c"]
    d_count = conn.execute("SELECT COUNT(*) AS c FROM dividend").fetchone()["c"]
    return {"holdings": h_count, "sales": s_count, "dividends": d_count}


def import_workbook(file_path_or_bytes: str | Path | io.BytesIO, replace: bool = False) -> dict[str, Any]:
    """Parse Invest.xlsx and populate database inside a single transaction.

    If data exists and replace=False, aborts without modification.
    Returns a detailed summary of imported records and any flagged manual values.
    """
    counts = check_existing_data()
    total_existing = sum(counts.values())

    if total_existing > 0 and not replace:
        return {
            "ok": False,
            "requires_confirmation": True,
            "message": (
                f"Existing data found ({counts['holdings']} holdings, {counts['sales']} sales, "
                f"{counts['dividends']} dividends). Please confirm with replace=True to overwrite."
            ),
            "existing": counts,
        }

    # Ensure raw bytes are available to open workbook twice (data_only=True and False)
    if isinstance(file_path_or_bytes, (str, Path)):
        raw_bytes = Path(file_path_or_bytes).read_bytes()
    elif isinstance(file_path_or_bytes, io.BytesIO):
        raw_bytes = file_path_or_bytes.getvalue()
    elif hasattr(file_path_or_bytes, "read"):
        raw_bytes = file_path_or_bytes.read()
    else:
        raw_bytes = bytes(file_path_or_bytes)

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    try:
        wb_formula = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=False)
    except Exception:
        wb_formula = None

    conn = get_connection()
    now = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")

    report: dict[str, Any] = {
        "ok": True,
        "imported_at": now,
        "portfolios": {},
        "flagged_items": [],
        "summary_values": 0,
    }

    conn.execute("BEGIN IMMEDIATE")
    try:
        # If replacing, clear existing Portfolio Tracker data
        if total_existing > 0 and replace:
            logger.info("Replacing existing Portfolio Tracker data as requested")
            conn.execute("DELETE FROM sale")
            conn.execute("DELETE FROM buy_lot")
            conn.execute("DELETE FROM holding")
            conn.execute("DELETE FROM dividend")
            conn.execute("DELETE FROM summary_value")

        # Map available sheets
        available_sheets = {name.strip().lower(): name for name in wb.sheetnames}

        for norm_port, target_sheet_name in SHEET_NAMES.items():
            lookup_key = target_sheet_name.strip().lower()
            if lookup_key not in available_sheets:
                logger.warning("Sheet %s not found in workbook — skipping", target_sheet_name)
                continue

            actual_sheet_name = available_sheets[lookup_key]
            ws = wb[actual_sheet_name]
            ws_f = wb_formula[actual_sheet_name] if (wb_formula and actual_sheet_name in wb_formula.sheetnames) else None
            layout = _resolve_layout(ws, norm_port)

            port_stats = {
                "open_holdings": 0,
                "buy_lots": 0,
                "sold_positions": 0,
                "dividends": 0,
            }

            # ---------------------------------------------------------------
            # 1. Parse Open Holdings
            # ---------------------------------------------------------------
            # Aggregation buffer for multi-buy rows: symbol -> {holding_meta, lots}
            open_holdings_buffer: dict[str, dict[str, Any]] = {}

            for r in range(layout["holdings_start"], ws.max_row + 1):
                col_a = ws.cell(r, 1).value
                if normalize_label(col_a) == TOTAL_ROW_LABEL:
                    break  # Reached the holdings total row
                if not col_a or not str(col_a).strip():
                    continue

                scheme = str(col_a).strip()

                inv_date = parse_sheet_date(ws.cell(r, 2).value)
                q, flag_q = _safe_number(ws.cell(r, 6).value)
                avg, flag_avg = _safe_number(ws.cell(r, 7).value)
                inv_val, _ = _safe_number(ws.cell(r, 9).value)
                buy_charge, _ = _safe_number(ws.cell(r, 10).value)
                app = str(ws.cell(r, 17).value or "").strip()
                remarks = str(ws.cell(r, 18).value or "").strip()

                if flag_q or flag_avg:
                    report["flagged_items"].append({
                        "portfolio": norm_port,
                        "section": "open_holding",
                        "row": r,
                        "scheme": scheme,
                        "note": f"Unusual quantity or avg price in Excel: Q={ws.cell(r,6).value!r}, Avg={ws.cell(r,7).value!r}",
                    })

                # Check manual Invested override (e.g. demerger with cost basis 0)
                formula_inv = ws_f.cell(r, 9).value if ws_f else None
                inv_is_formula = str(formula_inv or "").strip().startswith("=")
                expected_inv = q * avg
                if not inv_is_formula and abs(inv_val - expected_inv) > 0.01:
                    actual_inv = inv_val
                    report["flagged_items"].append({
                        "portfolio": norm_port,
                        "section": "open_holding",
                        "row": r,
                        "scheme": scheme,
                        "note": f"Manual Invested override preserved: {inv_val} (expected {expected_inv})",
                    })
                else:
                    actual_inv = None

                # Distinguish manual buy charge override vs derived value (§1.4)
                formula_val = ws_f.cell(r, 10).value if ws_f else None
                is_formula = str(formula_val or "").strip().startswith("=")
                expected_base = actual_inv if actual_inv is not None else expected_inv
                if _is_manual_charge(buy_charge, is_formula, compute_buy_charge(expected_base)):
                    actual_buy_charge = buy_charge
                    report["flagged_items"].append({
                        "portfolio": norm_port,
                        "section": "open_holding",
                        "row": r,
                        "scheme": scheme,
                        "note": f"Manual Buy Charge override preserved: {buy_charge}",
                    })
                else:
                    actual_buy_charge = None

                # Determine Person for LOAN portfolio
                person = map_app_to_person(app) if norm_port == "LOAN" else None

                lot_data = {
                    "invest_date": inv_date or "2025-01-01",
                    "quantity": q,
                    "avg_price": avg,
                    "invested_amount": actual_inv,
                    "buy_charge": actual_buy_charge,
                    "remarks": remarks,
                    "app": app,
                }

                if scheme not in open_holdings_buffer:
                    open_holdings_buffer[scheme] = {
                        "portfolio_name": norm_port,
                        "symbol": scheme,
                        "scheme_name": scheme,
                        "person": person,
                        "app": app,
                        "remarks": remarks,
                        "lots": [lot_data],
                    }
                else:
                    open_holdings_buffer[scheme]["lots"].append(lot_data)

            # Insert open holdings into database
            for scheme, h_data in open_holdings_buffer.items():
                cur = conn.execute(
                    """
                    INSERT INTO holding (
                        portfolio_name, symbol, scheme_name, name_confirmed,
                        person, app, remarks, status, created_at, updated_at
                    ) VALUES (?, ?, ?, 1, ?, ?, ?, 'open', ?, ?)
                    """,
                    (
                        h_data["portfolio_name"],
                        h_data["symbol"],
                        h_data["scheme_name"],
                        h_data["person"],
                        h_data["app"],
                        h_data["remarks"],
                        now,
                        now,
                    ),
                )
                holding_id = cur.lastrowid
                port_stats["open_holdings"] += 1

                for lot in h_data["lots"]:
                    conn.execute(
                        """
                        INSERT INTO buy_lot (
                            holding_id, invest_date, quantity, avg_price,
                            invested_amount, buy_charge, remarks, app, created_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            holding_id,
                            lot["invest_date"],
                            lot["quantity"],
                            lot["avg_price"],
                            lot.get("invested_amount"),
                            lot["buy_charge"],
                            lot["remarks"],
                            lot["app"],
                            now,
                        ),
                    )
                    port_stats["buy_lots"] += 1

            # ---------------------------------------------------------------
            # 2. Parse Sold Rows
            # ---------------------------------------------------------------
            sold_col = layout["sold_col"]

            for r in range(layout["sold_start"], ws.max_row + 1):
                col_val = ws.cell(r, sold_col).value
                if normalize_label(col_val) == TOTAL_ROW_LABEL:
                    break
                # Guard against entering dividend section in vertical stacked layout (MADI)
                if sold_col == 1 and (
                    normalize_label(col_val) in (DIVIDEND_SECTION_LABEL, "dividend", "dividends")
                    or r >= layout.get("dividend_start", 999999)
                ):
                    break
                if not col_val or not str(col_val).strip():
                    continue

                scheme = str(col_val).strip()
                inv_date = parse_sheet_date(ws.cell(r, sold_col + 1).value)
                sell_date = parse_sheet_date(ws.cell(r, sold_col + 2).value)
                q, flag_q = _safe_number(ws.cell(r, sold_col + 5).value)
                avg, flag_avg = _safe_number(ws.cell(r, sold_col + 6).value)
                sell_price, flag_sp = _safe_number(ws.cell(r, sold_col + 7).value)
                inv_val, _ = _safe_number(ws.cell(r, sold_col + 8).value)
                buy_charge, _ = _safe_number(ws.cell(r, sold_col + 9).value)
                sell_charge, _ = _safe_number(ws.cell(r, sold_col + 10).value)
                loss_val, _ = _safe_number(ws.cell(r, sold_col + 13).value)
                app = str(ws.cell(r, sold_col + 16).value or "").strip()
                remarks = str(ws.cell(r, sold_col + 17).value or "").strip()

                # Handle special advisory fee rows (e.g. Equity Premium, Mahveer Wealth with '-' as Q/Avg and -6999 as loss)
                if (flag_q or q == 0.0) and loss_val != 0.0:
                    q = 1.0
                    avg = abs(loss_val)
                    sell_price = 0.0
                    actual_inv = None
                    report["flagged_items"].append({
                        "portfolio": norm_port,
                        "section": "sold",
                        "row": r,
                        "scheme": scheme,
                        "note": f"Manual deduction/fee entry imported with loss {loss_val}: '{remarks or scheme}'",
                    })
                elif flag_q or flag_avg or flag_sp:
                    actual_inv = None
                    report["flagged_items"].append({
                        "portfolio": norm_port,
                        "section": "sold",
                        "row": r,
                        "scheme": scheme,
                        "note": f"Unusual value: Q={ws.cell(r,sold_col+5).value!r}, Avg={ws.cell(r,sold_col+6).value!r}, SellPrice={ws.cell(r,sold_col+7).value!r}",
                    })
                else:
                    # Check manual Invested override (e.g. demergers VEDPOWER, VOGL, VISL or partial lots)
                    formula_inv = ws_f.cell(r, sold_col + 8).value if ws_f else None
                    inv_is_formula = str(formula_inv or "").strip().startswith("=")
                    expected_inv = q * avg
                    if not inv_is_formula and abs(inv_val - expected_inv) > 0.01:
                        actual_inv = inv_val
                        report["flagged_items"].append({
                            "portfolio": norm_port,
                            "section": "sold",
                            "row": r,
                            "scheme": scheme,
                            "note": f"Manual Invested override preserved: {inv_val} (expected {expected_inv})",
                        })
                    else:
                        actual_inv = None

                person = map_app_to_person(app) if norm_port == "LOAN" else None

                # Distinguish manual overrides vs derived values (§1.4)
                formula_bc = ws_f.cell(r, sold_col + 9).value if ws_f else None
                formula_sc = ws_f.cell(r, sold_col + 10).value if ws_f else None
                bc_is_formula = str(formula_bc or "").strip().startswith("=")
                sc_is_formula = str(formula_sc or "").strip().startswith("=")
                expected_base = actual_inv if actual_inv is not None else (q * avg)
                actual_bc = (
                    buy_charge
                    if _is_manual_charge(buy_charge, bc_is_formula, compute_buy_charge(expected_base))
                    else None
                )
                actual_sc = (
                    sell_charge
                    if _is_manual_charge(sell_charge, sc_is_formula, compute_sell_charge(q * sell_price))
                    else None
                )

                if actual_bc is not None:
                    report["flagged_items"].append({
                        "portfolio": norm_port,
                        "section": "sold",
                        "row": r,
                        "scheme": scheme,
                        "note": f"Manual Buy Charge override preserved: {buy_charge}",
                    })
                if actual_sc is not None:
                    report["flagged_items"].append({
                        "portfolio": norm_port,
                        "section": "sold",
                        "row": r,
                        "scheme": scheme,
                        "note": f"Manual Sell Charge override preserved: {sell_charge}",
                    })

                # Insert holding with status='sold'
                h_cur = conn.execute(
                    """
                    INSERT INTO holding (
                        portfolio_name, symbol, scheme_name, name_confirmed,
                        person, app, remarks, status, created_at, updated_at
                    ) VALUES (?, ?, ?, 1, ?, ?, ?, 'sold', ?, ?)
                    """,
                    (norm_port, scheme, scheme, person, app, remarks, now, now),
                )
                holding_id = h_cur.lastrowid

                # Insert corresponding buy_lot
                conn.execute(
                    """
                    INSERT INTO buy_lot (
                        holding_id, invest_date, quantity, avg_price,
                        invested_amount, buy_charge, remarks, app, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        holding_id,
                        inv_date or "2024-01-01",
                        q if q > 0 else 1.0,
                        avg,
                        actual_inv,
                        actual_bc,
                        remarks,
                        app,
                        now,
                    ),
                )

                # Insert sale record
                conn.execute(
                    """
                    INSERT INTO sale (
                        holding_id, sell_date, quantity, sell_price,
                        sell_charge, remarks, app, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        holding_id,
                        sell_date or inv_date or "2024-01-01",
                        q if q > 0 else 1.0,
                        sell_price,
                        actual_sc,
                        remarks,
                        app,
                        now,
                    ),
                )
                port_stats["sold_positions"] += 1

            # ---------------------------------------------------------------
            # 3. Parse Dividends
            # ---------------------------------------------------------------
            for r in range(layout["dividend_start"], ws.max_row + 1):
                col_a = ws.cell(r, 1).value
                if normalize_label(col_a) == TOTAL_ROW_LABEL:
                    break
                if not col_a or not str(col_a).strip():
                    continue

                stock = str(col_a).strip()
                val, flag_val = _safe_number(ws.cell(r, 2).value)
                term = str(ws.cell(r, 3).value or "").strip()

                if flag_val:
                    report["flagged_items"].append({
                        "portfolio": norm_port,
                        "section": "dividend",
                        "row": r,
                        "scheme": stock,
                        "note": f"Unusual dividend amount: {ws.cell(r, 2).value!r}",
                    })

                if val > 0:
                    conn.execute(
                        """
                        INSERT INTO dividend (portfolio_name, symbol, value, received_date, created_at)
                        VALUES (?, ?, ?, ?, ?)
                        """,
                        (norm_port, stock, val, term or "2025-01-01", now),
                    )
                    port_stats["dividends"] += 1

            report["portfolios"][norm_port] = port_stats

        # -------------------------------------------------------------------
        # 4. Parse Summary Panel from Loan Sheet (Cols 39-40, Rows 76-87)
        # -------------------------------------------------------------------
        loan_sheet_name = available_sheets.get("loan", available_sheets.get("loan "))
        if loan_sheet_name and loan_sheet_name in wb.sheetnames:
            ws_loan = wb[loan_sheet_name]
            label_rows = {
                normalize_label(ws_loan.cell(r, SUMMARY_LABEL_COLUMN).value): r
                for r in range(1, ws_loan.max_row + 1)
                if ws_loan.cell(r, SUMMARY_LABEL_COLUMN).value
            }
            for item in SUMMARY_PANEL_CONFIG:
                if item["type"] != "fixed":
                    continue
                r = label_rows.get(normalize_label(item["label"]), item["row"])
                val, _ = _safe_number(ws_loan.cell(r, SUMMARY_VALUE_COLUMN).value)
                conn.execute(
                    """
                    INSERT INTO summary_value (key, value, label, updated_at)
                    VALUES (?, ?, ?, ?)
                    ON CONFLICT(key) DO UPDATE SET
                        value = excluded.value,
                        label = excluded.label,
                        updated_at = excluded.updated_at
                    """,
                    (item["key"], val, item["label"], now),
                )
                report["summary_values"] += 1

        conn.execute("COMMIT")
        logger.info("Successfully imported portfolio tracker data from workbook")
        return report

    except Exception as exc:
        conn.execute("ROLLBACK")
        logger.exception("Workbook import transaction failed and was rolled back: %s", exc)
        return {"ok": False, "error": str(exc), "message": "Import failed and changes were rolled back."}

# ---------------------------------------------------------------------------
# Exporter (§1.3)
# ---------------------------------------------------------------------------

def _write_holding_to_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    start_row: int,
    start_col: int,
    item: dict[str, Any],
) -> int:
    """Write an open holding to the sheet. If it has multiple buy lots, writes each lot as a row."""
    lots = item.get("lots") or []
    if len(lots) > 1:
        row = start_row
        for lot in lots:
            lot_item = {
                "symbol": item.get("symbol"),
                "scheme_name": item.get("scheme_name"),
                "first_invest_date": lot.get("invest_date"),
                "current_date": item.get("current_date"),
                "years": lot.get("years", 0),
                "months": lot.get("months", 0),
                "quantity": lot.get("quantity", 0),
                "avg_price": lot.get("avg_price", 0),
                "ltp": item.get("ltp", 0),
                "invested_amount": lot.get("invested_amount", 0),
                "buy_charge": lot.get("buy_charge", 0),
                "sell_charge": lot.get("sell_charge", 0),
                "current_total": lot.get("current_total", 0),
                "earned": lot.get("earned", 0),
                "loss": lot.get("loss", 0),
                "annual_return": lot.get("annual_return", 0),
                "total_return": lot.get("total_return", 0),
                "app": lot.get("app") or item.get("app"),
                "person": item.get("person"),
                "remarks": lot.get("remarks") or item.get("remarks"),
            }
            _write_holding_row_to_cells(ws, row, start_col, lot_item)
            row += 1
        return row
    else:
        _write_holding_row_to_cells(ws, start_row, start_col, item)
        return start_row + 1


def export_workbook(
    output_file_or_stream: str | Path | io.BytesIO,
    portfolio: str | None = None,
) -> None:
    """Generate an Excel workbook matching the exact format of Invest.xlsx.

    Can be scoped to a single portfolio or all three (§1.3).
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    total_font = Font(name="Calibri", size=11, bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    sold_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    ports_to_export = ["MADI", "BAPA", "LOAN"]
    if portfolio and portfolio.strip().upper() in ("MADI", "BAPA", "LOAN"):
        ports_to_export = [portfolio.strip().upper()]

    for port_key in ports_to_export:
        sheet_title = SHEET_NAMES[port_key]
        ws = wb.create_sheet(title=sheet_title)

        # Retrieve open holdings and enrich with child buy lots
        open_items = hold_repo.list_holdings(port_key, status="open")
        enriched_open = []
        for item in open_items:
            lots = hold_repo.get_lots(item["id"])
            enriched_open.append(enrich_holding_row(item, lots=lots))
        open_totals = compute_totals(enriched_open)

        sold_items = hold_repo.list_holdings(port_key, status="sold")
        enriched_sold = [enrich_sold_row(item) for item in sold_items]
        sold_totals = compute_totals(enriched_sold)

        dividends = div_repo.list_dividends(port_key)

        if port_key == "MADI":
            # Layout 1: Vertical Stack (Holdings -> Sold -> Dividends)
            ws.cell(1, 1, value="Portfolio Madi")
            ws.cell(1, 1).font = header_font

            # Row 2: Headers
            for c_idx, h in enumerate(HOLDING_COLUMNS, start=1):
                cell = ws.cell(2, c_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill

            # Rows 3+: Open Holdings (expanding multi-buy lots)
            curr_row = 3
            for item in enriched_open:
                curr_row = _write_holding_to_sheet(ws, curr_row, 1, item)

            # Holdings Total Row
            ws.cell(curr_row, 1, value="Total").font = total_font
            ws.cell(curr_row, 9, value=open_totals["invested_amount"]).font = total_font
            ws.cell(curr_row, 12, value=open_totals["current_total"]).font = total_font
            ws.cell(curr_row, 13, value=open_totals["earned"]).font = total_font
            ws.cell(curr_row, 14, value=open_totals["loss"]).font = total_font
            curr_row += 3

            # Sold Section Header
            ws.cell(curr_row, 1, value="Sold").font = header_font
            curr_row += 1

            # Sold Rows
            for item in enriched_sold:
                _write_sold_row_to_cells(ws, curr_row, 1, item)
                curr_row += 1

            # Sold Total Row
            ws.cell(curr_row, 1, value="Total").font = total_font
            ws.cell(curr_row, 9, value=sold_totals["invested_amount"]).font = total_font
            ws.cell(curr_row, 12, value=sold_totals["current_total"]).font = total_font
            ws.cell(curr_row, 13, value=sold_totals["earned"]).font = total_font
            ws.cell(curr_row, 14, value=sold_totals["loss"]).font = total_font
            curr_row += 2

            # Dividend Section
            ws.cell(curr_row, 1, value="Dividend").font = header_font
            curr_row += 1
            ws.cell(curr_row, 1, value="Stock").font = header_font
            ws.cell(curr_row, 2, value="Divi").font = header_font
            ws.cell(curr_row, 3, value="Term").font = header_font
            curr_row += 1

            div_sum = 0.0
            for d in dividends:
                ws.cell(curr_row, 1, value=d["symbol"])
                ws.cell(curr_row, 2, value=d["value"])
                ws.cell(curr_row, 3, value=d["received_date"])
                div_sum += d["value"]
                curr_row += 1

            ws.cell(curr_row, 1, value="Total").font = total_font
            ws.cell(curr_row, 2, value=div_sum).font = total_font

        else:
            # Layout 2: Side-by-Side (Bapa / Loan)
            ws.cell(1, 1, value="Stock Invest Details").font = header_font
            ws.cell(1, 20, value="Stock Sold").font = header_font

            # Row 2: Headers
            for c_idx, h in enumerate(HOLDING_COLUMNS, start=1):
                c1 = ws.cell(2, c_idx, value=h)
                c1.font = header_font
                c1.fill = header_fill
                c2 = ws.cell(2, c_idx + 19, value=h)
                c2.font = header_font
                c2.fill = sold_fill

            # Write Open Holdings in Cols 1-18 (expanding multi-buy lots)
            curr_row = 3
            for item in enriched_open:
                curr_row = _write_holding_to_sheet(ws, curr_row, 1, item)

            # Holdings Total
            ws.cell(curr_row, 1, value="Total").font = total_font
            ws.cell(curr_row, 9, value=open_totals["invested_amount"]).font = total_font
            ws.cell(curr_row, 12, value=open_totals["current_total"]).font = total_font
            ws.cell(curr_row, 13, value=open_totals["earned"]).font = total_font
            ws.cell(curr_row, 14, value=open_totals["loss"]).font = total_font

            # Write Sold in Cols 20-37
            sold_curr_row = 3
            for item in enriched_sold:
                _write_sold_row_to_cells(ws, sold_curr_row, 20, item)
                sold_curr_row += 1

            # Sold Total
            ws.cell(sold_curr_row, 20, value="Total").font = total_font
            ws.cell(sold_curr_row, 28, value=sold_totals["invested_amount"]).font = total_font
            ws.cell(sold_curr_row, 31, value=sold_totals["current_total"]).font = total_font
            ws.cell(sold_curr_row, 32, value=sold_totals["earned"]).font = total_font
            ws.cell(sold_curr_row, 33, value=sold_totals["loss"]).font = total_font

            # Dividends below Holdings
            div_start_row = max(curr_row + 3, 105)
            ws.cell(div_start_row, 1, value="Dividend").font = header_font
            div_start_row += 1
            ws.cell(div_start_row, 1, value="Stock").font = header_font
            ws.cell(div_start_row, 2, value="Divi").font = header_font
            ws.cell(div_start_row, 3, value="Term").font = header_font
            div_start_row += 1

            div_sum = 0.0
            for d in dividends:
                ws.cell(div_start_row, 1, value=d["symbol"])
                ws.cell(div_start_row, 2, value=d["value"])
                ws.cell(div_start_row, 3, value=d["received_date"])
                div_sum += d["value"]
                div_start_row += 1

            ws.cell(div_start_row, 1, value="Total").font = total_font
            ws.cell(div_start_row, 2, value=div_sum).font = total_font

            # Loan Summary Panel in Cols 39-40
            if port_key == "LOAN":
                sum_vals = sum_repo.get_all()
                loan_earned = open_totals["earned"] + sold_totals["earned"]
                loan_loss = open_totals["loss"] + sold_totals["loss"]
                panel = compute_summary_panel(
                    sum_vals,
                    loan_earned=loan_earned,
                    loan_loss=loan_loss,
                    loan_dividends=div_sum,
                )

                # Written from SUMMARY_PANEL_CONFIG so the labels the importer
                # matches on can never drift from the ones written here.
                panel_by_key = {
                    "current_stock_etf_invest": panel["block_a"]["current_stock_etf_invest"],
                    "current_mf_invest": panel["block_a"]["current_mf_invest"],
                    "current_mf_redeem": panel["block_a"]["current_mf_redeem"],
                    "_total_invest": panel["block_a"]["total"],
                    "loan_amount": panel["block_b"]["loan_amount"],
                    "_remaining_invest_loan": panel["block_b"]["remaining_invest_loan"],
                    "_stock_profit": panel["block_b"]["stock_profit"],
                    "_total_dividend": panel["block_b"]["dividend"],
                    "current_mf_redeem_profit": panel["block_b"]["current_mf_redeem_profit"],
                    "_current_remaining": panel["block_b"]["current_remaining"],
                }
                for item in SUMMARY_PANEL_CONFIG:
                    label_cell = ws.cell(item["row"], SUMMARY_LABEL_COLUMN, value=item["label"])
                    value_cell = ws.cell(
                        item["row"],
                        SUMMARY_VALUE_COLUMN,
                        value=panel_by_key.get(item["key"], 0.0),
                    )
                    if item["key"] in ("_total_invest", "_current_remaining"):
                        label_cell.font = total_font
                        value_cell.font = total_font

    if isinstance(output_file_or_stream, (str, Path)):
        wb.save(str(output_file_or_stream))
    else:
        wb.save(output_file_or_stream)


def export_csv(
    output_file_or_stream: Any,
    portfolio: str,
) -> None:
    """Export a single portfolio sheet as CSV (§1.3)."""
    import csv

    port_key = portfolio.strip().upper()
    if port_key not in ("MADI", "BAPA", "LOAN"):
        raise ValueError(f"Invalid portfolio: {portfolio}")

    # Retrieve open and sold
    open_items = hold_repo.list_holdings(port_key, status="open")
    enriched_open = []
    for item in open_items:
        lots = hold_repo.get_lots(item["id"])
        enriched_open.append(enrich_holding_row(item, lots=lots))
    open_totals = compute_totals(enriched_open)

    sold_items = hold_repo.list_holdings(port_key, status="sold")
    enriched_sold = [enrich_sold_row(item) for item in sold_items]
    sold_totals = compute_totals(enriched_sold)

    headers = [
        "Section", "Scheme", "Invest Date", "Current Date", "Years", "Months",
        "Quantity", "Avg Price", "LTP", "Invested Amount", "Buy Charge", "Sell Charge",
        "Current Total", "Earned", "Loss", "Annual Return %", "Total Return %",
        "App / Person", "Remarks"
    ]

    # Use a StringIO buffer if writing to bytes stream
    is_bytes = hasattr(output_file_or_stream, "write") and not hasattr(output_file_or_stream, "encoding")
    text_io = io.StringIO() if is_bytes else output_file_or_stream

    writer = csv.writer(text_io)
    writer.writerow([f"Portfolio Tracker — {port_key}"])
    writer.writerow([])
    writer.writerow(headers)

    # 1. Open holdings (expanding lots)
    for h in enriched_open:
        lots = h.get("lots") or []
        if len(lots) > 1:
            for lot in lots:
                writer.writerow([
                    "Open",
                    h.get("scheme_name") or h.get("symbol"),
                    format_sheet_date(lot.get("invest_date", "")),
                    format_sheet_date(h.get("current_date", "")),
                    lot.get("years", 0),
                    lot.get("months", 0),
                    lot.get("quantity", 0),
                    lot.get("avg_price", 0),
                    h.get("ltp", 0),
                    lot.get("invested_amount", 0),
                    lot.get("buy_charge", 0),
                    lot.get("sell_charge", 0),
                    lot.get("current_total", 0),
                    lot.get("earned", 0),
                    lot.get("loss", 0),
                    lot.get("annual_return", 0),
                    lot.get("total_return", 0),
                    lot.get("app") or h.get("app") or h.get("person") or "",
                    lot.get("remarks") or h.get("remarks", ""),
                ])
        else:
            writer.writerow([
                "Open",
                h.get("scheme_name") or h.get("symbol"),
                format_sheet_date(h.get("first_invest_date", "")),
                format_sheet_date(h.get("current_date", "")),
                h.get("years", 0),
                h.get("months", 0),
                h.get("quantity", 0),
                h.get("avg_price", 0),
                h.get("ltp", 0),
                h.get("invested_amount", 0),
                h.get("buy_charge", 0),
                h.get("sell_charge", 0),
                h.get("current_total", 0),
                h.get("earned", 0),
                h.get("loss", 0),
                h.get("annual_return", 0),
                h.get("total_return", 0),
                h.get("app") or h.get("person") or "",
                h.get("remarks", ""),
            ])

    writer.writerow([
        "Total Open", "", "", "", "", "", "", "", "",
        open_totals["invested_amount"], "", "", open_totals["current_total"],
        open_totals["earned"], open_totals["loss"], "", "", "", ""
    ])
    writer.writerow([])

    # 2. Sold positions
    writer.writerow(["--- Sold Positions ---"])
    writer.writerow(headers)
    for s in enriched_sold:
        writer.writerow([
            "Sold",
            s.get("scheme_name") or s.get("symbol"),
            format_sheet_date(s.get("invest_date", "")),
            format_sheet_date(s.get("sell_date", "")),
            s.get("years", 0),
            s.get("months", 0),
            s.get("quantity", 0),
            s.get("avg_price", 0),
            s.get("ltp", 0),
            s.get("invested_amount", 0),
            s.get("buy_charge", 0),
            s.get("sell_charge", 0),
            s.get("current_total", 0),
            s.get("earned", 0),
            s.get("loss", 0),
            s.get("annual_return", 0),
            s.get("total_return", 0),
            s.get("app") or s.get("person") or "",
            s.get("remarks", ""),
        ])

    writer.writerow([
        "Total Sold", "", "", "", "", "", "", "", "",
        sold_totals["invested_amount"], "", "", sold_totals["current_total"],
        sold_totals["earned"], sold_totals["loss"], "", "", "", ""
    ])

    if is_bytes:
        output_file_or_stream.write(text_io.getvalue().encode("utf-8"))


def _write_holding_row_to_cells(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, start_col: int, item: dict[str, Any]) -> None:
    """Helper to populate an 18-column open holding row."""
    ws.cell(row, start_col, value=item.get("symbol") or item.get("scheme_name"))
    ws.cell(row, start_col + 1, value=format_sheet_date(item.get("first_invest_date", "")))
    ws.cell(row, start_col + 2, value=format_sheet_date(item.get("current_date", "")))
    ws.cell(row, start_col + 3, value=item.get("years", 0))
    ws.cell(row, start_col + 4, value=item.get("months", 0))
    ws.cell(row, start_col + 5, value=item.get("quantity", 0))
    ws.cell(row, start_col + 6, value=item.get("avg_price", 0))
    ws.cell(row, start_col + 7, value=item.get("ltp", 0))
    ws.cell(row, start_col + 8, value=item.get("invested_amount", 0))
    ws.cell(row, start_col + 9, value=item.get("buy_charge", 0))
    ws.cell(row, start_col + 10, value=item.get("sell_charge", 0))
    ws.cell(row, start_col + 11, value=item.get("current_total", 0))
    ws.cell(row, start_col + 12, value=item.get("earned", 0))
    ws.cell(row, start_col + 13, value=item.get("loss", 0))
    ws.cell(row, start_col + 14, value=item.get("annual_return", 0))
    ws.cell(row, start_col + 15, value=item.get("total_return", 0))
    ws.cell(row, start_col + 16, value=item.get("app") or item.get("person") or "")
    ws.cell(row, start_col + 17, value=item.get("remarks", ""))


def _write_sold_row_to_cells(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, start_col: int, item: dict[str, Any]) -> None:
    """Helper to populate an 18-column sold row."""
    ws.cell(row, start_col, value=item.get("symbol") or item.get("scheme_name"))
    ws.cell(row, start_col + 1, value=format_sheet_date(item.get("invest_date", "")))
    ws.cell(row, start_col + 2, value=format_sheet_date(item.get("sell_date", "")))
    ws.cell(row, start_col + 3, value=item.get("years", 0))
    ws.cell(row, start_col + 4, value=item.get("months", 0))
    ws.cell(row, start_col + 5, value=item.get("quantity", 0))
    ws.cell(row, start_col + 6, value=item.get("avg_price", 0))
    ws.cell(row, start_col + 7, value=item.get("ltp", 0))
    ws.cell(row, start_col + 8, value=item.get("invested_amount", 0))
    ws.cell(row, start_col + 9, value=item.get("buy_charge", 0))
    ws.cell(row, start_col + 10, value=item.get("sell_charge", 0))
    ws.cell(row, start_col + 11, value=item.get("current_total", 0))
    ws.cell(row, start_col + 12, value=item.get("earned", 0))
    ws.cell(row, start_col + 13, value=item.get("loss", 0))
    ws.cell(row, start_col + 14, value=item.get("annual_return", 0))
    ws.cell(row, start_col + 15, value=item.get("total_return", 0))
    ws.cell(row, start_col + 16, value=item.get("app") or item.get("person") or "")
    ws.cell(row, start_col + 17, value=item.get("remarks", ""))
